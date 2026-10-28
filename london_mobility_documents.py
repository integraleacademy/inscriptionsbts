"""The original London Word model and its compiled, runtime-independent PDF.

The PDF background is exported from the same model, with reserved areas for the
variable paragraphs. No Office installation or network service is used online.
"""
from datetime import date
from io import BytesIO
import hashlib
import json
from pathlib import Path
import re
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

from certificates import _replace
from person_names import person_name

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / 'assets' / 'london_mobility'
MODEL = ASSETS / 'convention_mobilite.dotx'
VERSION = 'london-20261028-1'


def date_fr(value):
    return date.fromisoformat(value).strftime('%d/%m/%Y')


def paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from paragraphs(cell)


def minor_at_signing(values):
    born = date.fromisoformat(values['apprentice_birth_date'])
    signing = date.fromisoformat(values['convention_date'])
    return (signing.year - born.year - ((signing.month, signing.day) < (born.month, born.day))) < 18


def variable_text(values):
    days = (date.fromisoformat(values['mobility_end']) - date.fromisoformat(values['mobility_start'])).days + 1
    name = person_name(values['apprentice_last_name'], values['apprentice_first_name'])
    civility = {'Homme': 'Monsieur', 'Femme': 'Madame', 'M': 'Monsieur', 'F': 'Madame'}.get(values.get('apprentice_sex'), '')
    beneficiary = ' '.join(filter(None, (civility, name)))
    return {
        'company': 'Raison sociale : ' + values['employer_name'],
        'address': 'Adresse : ' + values['employer_address'],
        'city': 'Code postal : ' + values['employer_postcode'] + ' – Ville : ' + values['employer_city'],
        'siret': 'SIRET : ' + values['employer_siret'],
        'beneficiary': 'Nom et prénom(s) : ' + beneficiary,
        'period': f"La présente convention s’applique du {date_fr(values['mobility_start'])} au {date_fr(values['mobility_end'])}, soit une durée totale de {days} jours.",
        'issue': 'Fait à Puget sur Argens, le ' + date_fr(values['convention_date']) + '.',
        'company_signature': values['employer_name'],
        'beneficiary_signature': beneficiary,
        'guardian': 'Pour le représentant légal : ' + person_name(values.get('guardian_last_name'), values.get('guardian_first_name')),
    }


def word_document(values, *, compile_fields=False):
    """Keep source clauses, annexes, graphics and run styles intact."""
    output = BytesIO()
    with zipfile.ZipFile(MODEL) as source, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
        for name in source.namelist():
            data = source.read(name)
            if name == '[Content_Types].xml':
                data = data.replace(b'wordprocessingml.template.main+xml', b'wordprocessingml.document.main+xml')
            if name.startswith('word/') and name.endswith('.xml'):
                root = etree.fromstring(data)
                # The source mail-merge instructions must not resurrect fields.
                for tag in ('instrText', 'fldChar'):
                    for element in list(root.iter(qn('w:' + tag))):
                        element.getparent().remove(element)
                data = etree.tostring(root, encoding='utf-8', xml_declaration=True)
            target.writestr(name, data)
    output.seek(0)
    doc = Document(output)
    texts = variable_text(values)
    for paragraph in paragraphs(doc):
        content = paragraph.text.strip()
        field = None
        if '«NOM_ENTREPRISE»' in content:
            field = 'company' if content.startswith('Raison') else 'company_signature'
        elif '«ADR1_ENTREPRISE»' in content: field = 'address'
        elif '«CP_ENTREPRISE»' in content: field = 'city'
        elif '«SIRET_ENTREPRISE»' in content: field = 'siret'
        elif '«NOM_PRENOM_APPRENANT»' in content:
            field = 'beneficiary' if content.startswith('Nom') else 'beneficiary_signature'
        elif content.startswith("La présente convention s'applique du"): field = 'period'
        elif content.startswith('Fait en 4 exemplaires'): field = 'issue'
        if field:
            # Keep the first run's existing typography and the paragraph style.
            _replace(paragraph._p, {paragraph.text: f'LONDONFIELD_{field.upper()}' if compile_fields else texts[field]})
            if compile_fields:
                # Reserve two lines for long identities/addresses and duration.
                if field in {'company', 'address', 'city', 'beneficiary', 'company_signature', 'beneficiary_signature', 'period'}:
                    paragraph.add_run('\n ')
        if '[sc_' in paragraph.text or '[/sc_' in paragraph.text:
            content = paragraph.text
            _replace(paragraph._p, {m.group(): '' for m in re.finditer(r'\[/?sc_[^\]]+\]', content)})
            if '[/sc_sign1.signature]' in content and minor_at_signing(values):
                # Use the source's existing blank signature area, without
                # pushing the host signature onto an otherwise empty page.
                guardian = paragraph
                guardian.add_run('LONDONFIELD_GUARDIAN' if compile_fields else texts['guardian']).font.size = Pt(11)
                if compile_fields:
                    guardian.add_run('\n ')
    for paragraph in paragraphs(doc):
        if re.search(r'«[A-Z_0-9]+»|\[/?sc_', paragraph.text):
            raise ValueError('Une variable du modèle de mobilité est restée vide.')
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def pdf_document(values):
    from pypdf import PdfReader, PdfWriter
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas

    manifest = json.loads((ASSETS / 'manifest.json').read_text())
    if hashlib.sha256(MODEL.read_bytes()).hexdigest() != manifest['source_sha256']:
        raise RuntimeError('Le modèle de convention a changé. Son export PDF doit être actualisé.')
    variant = manifest['minor' if minor_at_signing(values) else 'adult']
    reader, writer = PdfReader(ASSETS / variant['file']), PdfWriter()
    font = 'LondonMobility'
    if font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font, str(ROOT / 'assets/certificates/fonts/LiberationSans-Regular.ttf')))
    texts = variable_text(values)
    for index, page in enumerate(reader.pages):
        packet = BytesIO()
        drawing = Canvas(packet, pagesize=(float(page.mediabox.width), float(page.mediabox.height)))
        for field in variant['fields']:
            if field['page'] != index: continue
            text, size = texts[field['name']], field['size']
            width, max_lines = field['width'], field['lines']
            while size >= 9:
                lines, line = [], ''
                for word in text.split():
                    proposed = (line + ' ' + word).strip()
                    if line and pdfmetrics.stringWidth(proposed, font, size) > width:
                        lines.append(line); line = word
                    else: line = proposed
                if line: lines.append(line)
                if len(lines) <= max_lines and all(pdfmetrics.stringWidth(s, font, size) <= width for s in lines): break
                size -= .5
            else:
                raise ValueError('Le texte « ' + texts[field['name']][:40] + '… » est trop long pour le modèle. Abrégez ce champ.')
            drawing.setFont(font, size)
            for offset, line in enumerate(lines):
                y = float(page.mediabox.height) - field['baseline'] - offset * field['leading']
                if field.get('align') == 'center':
                    drawing.drawCentredString(field['x'] + field['width']/2, y, line)
                else:
                    drawing.drawString(field['x'], y, line)
        drawing.showPage(); drawing.save(); packet.seek(0)
        page.merge_page(PdfReader(packet).pages[0])
        writer.add_page(page)
    writer.add_metadata({'/Title': 'Convention de mobilité internationale – Londres', '/Author': 'Intégrale Academy'})
    out = BytesIO(); writer.write(out)
    return out.getvalue()
