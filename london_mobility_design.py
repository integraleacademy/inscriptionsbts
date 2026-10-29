"""Six-page London layout. Legal text and existing seals come from the source model."""
from functools import lru_cache
from html import escape
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from person_names import person_name

ROOT = Path(__file__).resolve().parent
MODEL = ROOT / 'assets/london_mobility/convention_mobilite.dotx'
LOGO = ROOT / 'assets/student_cards/integrale_academy_official.png'


@lru_cache(maxsize=1)
def source():
    with ZipFile(MODEL) as archive:
        ns = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        body = etree.fromstring(archive.read('word/document.xml')).find('w:body', ns)
        texts = [''.join(p.xpath('.//w:t/text()',namespaces=ns)).strip() for p in body.findall('w:p',ns)]
        assert texts[57].startswith('Article 1') and texts[241].startswith('La période de mobilité')
        images = {i:archive.read('word/media/image%d.png' % i) for i in (2,3,4)}
    return texts, images


def pages(values):
    from london_mobility_documents import date_fr, variable_text
    p, _ = source(); v = variable_text(values)
    def article(title, indexes):
        return [('h2',p[title])] + [('p',v['period'] if i == 72 else p[i]) for i in indexes]
    page2 = [('h1','Le cadre de la mobilité')] + article(57,[59,61,63,65,67]) + article(70,[72]) + article(76,[78,80,82,84])
    page3 = [('h1','Les engagements des parties')] + article(87,[89,91]) + article(93,[95,96,98,100,101,103]) + article(107,[109,111,113,115,117]) + article(120,[122,123])
    page5 = [('kicker','ANNEXE 1 · ANNEXE PÉDAGOGIQUE'),('h1','Objectifs et programme'),('h2',p[179]),('small',p[180]),('p',p[182]),('p',p[184])] + [('bullet',p[i]) for i in range(186,197)] + [('h2','Programme de la période de mobilité'),('p',p[198])] + [('bullet',p[i]) for i in range(200,210)]
    page6 = [('kicker','ANNEXE 1 · ANNEXE PÉDAGOGIQUE'),('h1','Activités et évaluation'),('h2',p[210]),('p',p[212]),('p',p[214]),('p',p[216])] + [('bullet',p[i]) for i in range(218,226)] + [('h2',p[226]),('p',p[228])] + [('bullet',p[i]) for i in range(230,235)] + [('h2',p[235])] + [('p',p[i]) for i in [237,239,241]]
    return {2:page2,3:page3,5:page5,6:page6}


def parties(values):
    from london_mobility_documents import variable_text, date_fr
    p,_ = source(); v = variable_text(values)
    return [
        ('L’employeur français',[v['company'],v['address'],v['city'],v['siret'],p[14]]),
        ('Le bénéficiaire du contrat d’apprentissage',[v['beneficiary'],'Date de naissance : '+date_fr(values['apprentice_birth_date']),p[53]]),
        ('L’organisme de formation et CFA français',[p[i] for i in [18,19,20,21,22,23,24,25,27]]),
        ('L’organisme de formation et CFA d’accueil',[p[i] for i in [40,41,42,43,44,45,46,48]]),
    ]


def signature_layout(values):
    from london_mobility_documents import minor_at_signing
    fields = {'employer':dict(type='signature',page=4,x=48,y=330,width=220,height=50),
              'apprentice':dict(type='signature',page=4,x=48,y=555,width=220,height=50)}
    if minor_at_signing(values): fields['guardian'] = dict(type='signature',page=4,x=48,y=735,width=220,height=45)
    return fields


def pdf_document(values):
    from reportlab.lib.colors import HexColor
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Paragraph
    from london_mobility_documents import date_fr, variable_text, minor_at_signing
    for suffix, file in [('', 'Regular'),('B','Bold')]:
        name = 'LondonModern'+suffix
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name,str(ROOT/'assets/certificates/fonts'/('LiberationSans-'+file+'.ttf'))))
    pdfmetrics.registerFontFamily('LondonModern',normal='LondonModern',bold='LondonModernB')
    out=BytesIO(); width,height=595.2756,841.8898
    c=Canvas(out,pagesize=(width,height),pageCompression=1)
    c.setTitle('Convention de mobilité internationale à Londres'); c.setAuthor('Intégrale Academy')
    ink,muted,line = '#252B34','#68788A','#DFE5ED'
    styles = {
        'p':ParagraphStyle('body',fontName='LondonModern',fontSize=10.5,leading=14.2,textColor=HexColor(ink),spaceAfter=7),
        'small':ParagraphStyle('small',fontName='LondonModern',fontSize=8.5,leading=11.5,textColor=HexColor(muted),spaceAfter=7),
        'h1':ParagraphStyle('title',fontName='LondonModernB',fontSize=21,leading=26,textColor=HexColor('#000000'),spaceAfter=20),
        'h2':ParagraphStyle('heading',fontName='LondonModernB',fontSize=11.5,leading=15,textColor=HexColor('#000000'),spaceAfter=8,spaceBefore=8),
        'kicker':ParagraphStyle('kicker',fontName='LondonModernB',fontSize=8,leading=11,textColor=HexColor(muted),spaceAfter=10),
        'bullet':ParagraphStyle('bullet',fontName='LondonModern',fontSize=10.5,leading=14.2,textColor=HexColor(ink),spaceAfter=3,leftIndent=12,firstLineIndent=-10),
    }
    def text(value,x,y,w,kind='p'):
        style=styles[kind]; y += style.spaceBefore
        paragraph=Paragraph(('• ' if kind=='bullet' else '')+escape(value),style)
        _,h=paragraph.wrap(w,800)
        if y+h>790: raise ValueError('Le contenu de la convention dépasse la zone imprimable. Vérifiez la longueur des informations sources.')
        paragraph.drawOn(c,x,height-y-h)
        return y+h+style.spaceAfter
    def header(number,label):
        c.drawImage(str(LOGO),48,height-64,width=33,height=33,mask='auto')
        c.setFont('LondonModernB',10); c.setFillColor(HexColor('#202D43')); c.drawString(91,height-43,'INTÉGRALE ACADEMY')
        c.setFont('LondonModern',8);c.setFillColor(HexColor(muted));c.drawString(91,height-57,'Mobilité internationale · Londres')
        c.setFont('LondonModernB',8);c.drawRightString(width-48,height-44,label.upper())
        c.setStrokeColor(HexColor(line));c.setLineWidth(.6);c.line(48,40,width-48,40)
        c.setFont('LondonModern',8);c.drawString(48,25,'Convention de mobilité · '+date_fr(values['mobility_start'])+' au '+date_fr(values['mobility_end']))
        c.drawRightString(width-48,25,f'{number} / 6')
    p,images=source();v=variable_text(values)
    header(1,'Les parties')
    y=text('Convention de mobilité internationale',48,90,500,'h1')
    y=text('Londres · Royaume-Uni',48,y-9,500,'h2')
    y=text(p[1],48,y,499,'p')
    y=text(p[2],48,y,499,'small')
    y=text('Du '+date_fr(values['mobility_start'])+' au '+date_fr(values['mobility_end']),48,y+5,499,'h2')
    records = parties(values)
    for pair in (records[:2],records[2:]):
        bottoms = []
        for x,(heading,lines) in zip((48,318),pair):
            bottom=text(heading,x,y+9,229,'h2')
            for value in lines:
                bottom=text(value,x,bottom,229,'small' if value.startswith('ci-après') else 'p') - 4
            bottoms.append(bottom)
        y=max(bottoms)+16
    if y>790: raise ValueError('Les informations des parties sont trop longues pour la convention.')
    c.showPage()
    content=pages(values)
    for number in range(2,7):
        header(number, 'Signatures' if number==4 else 'Annexe pédagogique' if number>4 else 'Convention')
        if number!=4:
            y=90
            for kind,value in content[number]: y=text(value,48,y,499,kind)
        else:
            y=text('Signatures des parties',48,90,499,'h1')
            text(v['issue'],48,y,499)
            text('Convention de mobilité du '+date_fr(values['mobility_start'])+' au '+date_fr(values['mobility_end']),48,y+28,499,'small')
            slots=[(48,205,'Pour l’employeur français',values['employer_name']),
                   (318,205,'Pour l’organisme de formation et CFA français','Clément VAILLANT · Directeur général'),
                   (48,440,'Pour le bénéficiaire du contrat d’apprentissage',person_name(values['apprentice_last_name'],values['apprentice_first_name'])),
                   (318,440,'Pour l’organisme de formation et CFA d’accueil','KAPLAN INTERNATIONAL · Rudy JARMUSZCZAK · Directeur France')]
            for x,y,label,name in slots:
                y=text(label,x,y,229,'h2');text(name,x,y,229)
            c.drawImage(ImageReader(BytesIO(images[2])),318,height-338,width=105,height=59,mask='auto')
            c.drawImage(ImageReader(BytesIO(images[3])),318,height-402,width=190,height=63,mask='auto')
            c.drawImage(ImageReader(BytesIO(images[4])),318,height-605,width=200,height=61,mask='auto')
            if minor_at_signing(values):
                y=text('Pour le représentant légal',48,650,499,'h2')
                text(person_name(values['guardian_last_name'],values['guardian_first_name']),48,y,499)
            for field in signature_layout(values).values():
                text('Signature électronique',field['x'],field['y']-18,field['width'],'small')
                c.setStrokeColor(HexColor(line));c.setLineWidth(.7)
                c.roundRect(field['x'],height-field['y']-field['height'],field['width'],field['height'],5,stroke=1,fill=0)
        c.showPage()
    c.save();return out.getvalue()


def word_document(values):
    from docx import Document
    from docx.shared import Pt, Mm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from london_mobility_documents import date_fr, variable_text, minor_at_signing
    doc=Document()
    for border in list(doc.styles.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    s=doc.sections[0]
    s.page_width=Mm(210);s.page_height=Mm(297)
    s.top_margin=Mm(30);s.bottom_margin=Mm(24);s.left_margin=s.right_margin=Mm(17)
    s.header_distance=Mm(10);s.footer_distance=Mm(13)
    for name,size in [('Normal',10.5),('Title',22),('Heading 1',20),('Heading 2',11.5)]:
        style=doc.styles[name];style.font.name='Liberation Sans';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_after=Pt(7);style.paragraph_format.line_spacing=1.12
        if name.startswith('Heading'):style.paragraph_format.space_before=Pt(9);style.font.bold=True
    doc.styles['Normal'].paragraph_format.widow_control=True
    h=s.header.paragraphs[0];h.paragraph_format.space_after=Pt(0)
    h.add_run().add_picture(str(LOGO),width=Mm(11));r=h.add_run('   INTÉGRALE ACADEMY   /   MOBILITÉ INTERNATIONALE');r.font.size=Pt(9);r.bold=True
    f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=f.add_run('Convention de mobilité · Londres     |     ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('68788A')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
    def add(kind,value):
        style={'h1':'Heading 1','h2':'Heading 2'}.get(kind,'Normal')
        p=doc.add_paragraph(style=style)
        p.add_run(('• ' if kind=='bullet' else '')+value)
        if kind in {'small','kicker'}:
            for r in p.runs:r.font.size=Pt(8.5);r.font.color.rgb=RGBColor.from_string('68788A')
        if kind=='bullet':
            p.paragraph_format.left_indent=Mm(4);p.paragraph_format.first_line_indent=Mm(-3);p.paragraph_format.space_after=Pt(3)
        return p
    p,images=source();v=variable_text(values)
    doc.add_paragraph('Convention de mobilité internationale',style='Title')
    add('h2','Londres · Royaume-Uni');add('p',p[1]);add('small',p[2])
    add('h2','Du '+date_fr(values['mobility_start'])+' au '+date_fr(values['mobility_end']))
    for heading,lines in parties(values):
        add('h2',heading)
        for value in lines:
            para=add('small' if value.startswith('ci-après') else 'p',value)
            para.paragraph_format.space_after=Pt(2)
    content=pages(values)
    for number in range(2,7):
        doc.add_page_break()
        if number!=4:
            for kind,value in content[number]:add(kind,value)
            continue
        add('h1','Signatures des parties');add('p',v['issue'])
        add('small','Convention de mobilité du '+date_fr(values['mobility_start'])+' au '+date_fr(values['mobility_end']))
        table=doc.add_table(rows=2,cols=2);table.autofit=False
        for col in table.columns:col.width=Mm(88)
        rows=[('Pour l’employeur français',values['employer_name'],None),
              ('Pour l’organisme de formation et CFA français','Clément VAILLANT · Directeur général',[2,3]),
              ('Pour le bénéficiaire du contrat d’apprentissage',person_name(values['apprentice_last_name'],values['apprentice_first_name']),None),
              ('Pour l’organisme de formation et CFA d’accueil','KAPLAN INTERNATIONAL · Rudy JARMUSZCZAK · Directeur France',[4])]
        for cell,(label,name,graphics) in zip([c for row in table.rows for c in row.cells],rows):
            cell.paragraphs[0].add_run(label).bold=True;cell.add_paragraph(name)
            if graphics:
                for i in graphics:cell.add_paragraph().add_run().add_picture(BytesIO(images[i]),width=Mm(35 if i==2 else 65))
            else:
                cell.add_paragraph('\n\n\nSignature').paragraph_format.space_after=Pt(30)
            props=cell._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            props.append(borders);margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'170');e.set(qn('w:type'),'dxa');margins.append(e)
            props.append(margins)
        if minor_at_signing(values):
            add('h2','Pour le représentant légal');add('p',person_name(values['guardian_last_name'],values['guardian_first_name']));add('p','Signature :')
    doc.core_properties.title='Convention de mobilité internationale à Londres';doc.core_properties.author='Intégrale Academy'
    out=BytesIO();doc.save(out);return out.getvalue()
