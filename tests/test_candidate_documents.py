"""Admin document viewing: no live data or external I/O."""
from io import BytesIO
import json
from pathlib import Path
import sqlite3

from docx import Document
from flask import Flask, session
from PIL import Image
import pytest

from candidate_documents import register_candidate_documents, document_preview_response


@pytest.fixture
def docs(tmp_path):
    root = Path(__file__).resolve().parents[1]
    app = Flask(__name__, template_folder=str(root / 'templates'))
    app.secret_key = 'document-tests-only'
    app.config['TESTING'] = True
    upload = tmp_path / 'uploads'
    upload.mkdir()
    (upload / 'alice').mkdir()
    (upload / 'bob').mkdir()
    database = tmp_path / 'test.db'

    def db():
        conn = sqlite3.connect(database)
        conn.row_factory = sqlite3.Row
        return conn

    with db() as conn:
        conn.execute('CREATE TABLE candidats(id TEXT PRIMARY KEY, fichiers_ci TEXT, fichiers_photo TEXT)')
        conn.execute('INSERT INTO candidats VALUES(?,?,?)', ('alice', '[]', '[]'))
        conn.execute('INSERT INTO candidats VALUES(?,?,?)', ('bob', '[]', '[]'))
    register_candidate_documents(app, db, lambda: str(upload), lambda: session.get('admin_ok'))
    return app.test_client(), db, upload, app


def login(client):
    with client.session_transaction() as state:
        state['admin_ok'] = True


def test_pdf_is_inline_private_and_requires_admin(docs):
    client, _, upload, _ = docs
    (upload / 'alice' / 'CV_alice.pdf').write_bytes(b'%PDF-1.4\nfixture')
    url = '/admin/candidat/alice/document/CV_alice.pdf'
    assert client.get(url).status_code == 403
    login(client)
    response = client.get(url)
    assert response.status_code == 200
    assert response.mimetype == 'application/pdf'
    assert response.headers['Content-Disposition'].startswith('inline;')
    assert response.headers['Cache-Control'] == 'private, no-store'
    assert response.headers['X-Content-Type-Options'] == 'nosniff'
    assert client.get(url + '?download=1').headers['Content-Disposition'].startswith('attachment;')


def test_file_cannot_be_taken_from_another_dossier_or_symlink(docs):
    client, db, upload, _ = docs
    login(client)
    foreign = upload / 'bob' / 'CV_bob.pdf'
    foreign.write_bytes(b'%PDF-1.4\nfixture')
    (upload / 'alice' / 'borrowed.pdf').symlink_to(foreign)
    with db() as conn:
        conn.execute('UPDATE candidats SET fichiers_ci=? WHERE id=?', (json.dumps([str(foreign)]), 'alice'))
    assert client.get('/admin/candidat/alice/document/CV_bob.pdf').status_code == 404
    assert client.get('/admin/candidat/alice/document/borrowed.pdf').status_code == 404
    assert client.get('/admin/candidat/unknown/document/CV_bob.pdf').status_code == 404
    assert client.get('/admin/candidat/alice/document/..%2Fbob%2FCV_bob.pdf').status_code == 404


def test_old_prefix_and_flat_upload_require_exact_dossier(docs):
    client, db, upload, _ = docs
    login(client)
    (upload / 'CI_alice.pdf').write_bytes(b'%PDF-1.4\nfixture')
    (upload / 'CI_bob.pdf').write_bytes(b'%PDF-1.4\nfixture')
    with db() as conn:
        conn.execute('UPDATE candidats SET fichiers_ci=? WHERE id=?', (json.dumps(['/old/prefix/CI_alice.pdf', '/old/prefix/CI_bob.pdf']), 'alice'))
    assert client.get('/admin/candidat/alice/document/CI_alice.pdf').status_code == 200
    assert client.get('/admin/candidat/alice/document/CI_bob.pdf').status_code == 404
    assert client.get('/admin/candidat/bob/document/CI_bob.pdf').status_code == 404


def test_image_type_is_detected_from_content(docs):
    client, _, upload, _ = docs
    login(client)
    Image.new('RGB', (8, 8), 'red').save(upload / 'alice' / 'picture.bin', 'PNG')
    response = client.get('/admin/candidat/alice/document/picture.bin')
    assert response.status_code == 200 and response.mimetype == 'image/png'
    assert response.headers['Content-Disposition'].startswith('inline;')


def test_html_disguised_as_pdf_or_svg_is_never_executed(docs):
    client, _, upload, _ = docs
    login(client)
    for name in ('fake.pdf', 'evil.svg'):
        (upload / 'alice' / name).write_text('<script>alert(document.cookie)</script>')
        response = client.get('/admin/candidat/alice/document/' + name)
        assert response.mimetype == 'text/html'
        assert b'<script>' not in response.data
        assert 'default-src' in response.headers['Content-Security-Policy']
        assert 'Télécharger l’original' in response.text


def test_docx_preview_escapes_text_and_preserves_tables(docs):
    client, _, upload, _ = docs
    login(client)
    doc = Document()
    doc.add_paragraph('<script>alert(1)</script>')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Promotion 2027'
    table.cell(0, 1).text = 'Une & deux'
    picture = BytesIO()
    Image.new('RGB', (20, 20), 'green').save(picture, 'PNG')
    picture.seek(0)
    doc.add_picture(picture)
    doc.save(upload / 'alice' / 'certificate.docx')
    response = client.get('/admin/candidat/alice/document/certificate.docx')
    assert response.status_code == 200 and response.mimetype == 'text/html'
    assert '<script>' not in response.text and '&lt;script&gt;' in response.text
    assert 'Promotion 2027' in response.text and '<table>' in response.text
    assert 'data:image/jpeg;base64,' in response.text
    assert '?download=1' in response.text


def test_certificate_helper_keeps_explicit_download_link(docs):
    _, _, upload, app = docs
    doc = Document()
    doc.add_paragraph('Certificat de scolarité 2027–2029')
    path = upload / 'certificate.docx'
    doc.save(path)
    with app.test_request_context('/admin/generate_certificat/alice?preview=1'):
        response = document_preview_response(path, title='Certificat de scolarité', download_url='/admin/generate_certificat/alice')
        assert response.mimetype == 'text/html'
        assert 'Certificat de scolarité 2027–2029' in response.get_data(as_text=True)
        assert 'href="/admin/generate_certificat/alice"' in response.get_data(as_text=True)


def test_invalid_docx_opens_fallback_page_without_auto_download(docs):
    client, _, upload, _ = docs
    login(client)
    (upload / 'alice' / 'broken.docx').write_text('not a zip')
    response = client.get('/admin/candidat/alice/document/broken.docx')
    assert response.status_code == 200
    assert 'aperçu de ce document n’est pas disponible' in response.text


def test_candidate_directory_cannot_alias_another_candidate(docs):
    client, _, upload, _ = docs
    login(client)
    (upload / 'bob' / 'CV_bob.pdf').write_bytes(b'%PDF-1.4\nfixture')
    (upload / 'alice').rmdir()
    (upload / 'alice').symlink_to(upload / 'bob', target_is_directory=True)
    assert client.get('/admin/candidat/alice/document/CV_bob.pdf').status_code == 404
