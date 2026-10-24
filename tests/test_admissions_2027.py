"""Campaign boundary, cohort isolation and submission regressions (no external sends)."""
import io
import json
import os
import tempfile
from datetime import timedelta
from unittest.mock import patch

os.environ.setdefault('INTEGRITY_CHECK', '0')
os.environ.setdefault('DATA_DIR', tempfile.mkdtemp(prefix='bts-campaign-bootstrap-'))
import pytest
import app as application
from campaign import ADMISSION_YEAR, OPENS_AT, candidate_year, campaign_is_open
from mail_templates import mail_html
from utils import dossier_number, get_mail_context


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(application, 'DATA_DIR', str(tmp_path))
    monkeypatch.setattr(application, 'DB_PATH', str(tmp_path / 'app.db'))
    monkeypatch.setattr(application, 'UPLOAD_DIR', str(tmp_path / 'uploads'))
    monkeypatch.setattr(application, 'PORTAL_FILE', str(tmp_path / 'portal.json'))
    monkeypatch.setattr(application, 'send_mail', lambda *args, **kwargs: True)
    monkeypatch.setattr(application, 'send_sms_brevo', lambda *args, **kwargs: True)
    application.init_db()
    application.ensure_schema()
    application.ensure_apprentissage_fields()
    application.ensure_baccalaureat_field()
    application.ensure_logs_cid_column()
    application.ensure_relance_field()
    application.ensure_souhaite_accompagnement_field()
    application.ensure_campaign_schema()
    return application.app.test_client()


def auth(client):
    with client.session_transaction() as session:
        session['admin_ok'] = True


def seed(year, cid, status='confirmee'):
    with application.db() as conn:
        conn.execute('INSERT INTO candidats(id,nom,prenom,email,bts,mode,numero_dossier,rentree_annee,statut,created_at,slug_public) VALUES(?,?,?,?,?,?,?,?,?,?,?)', (cid,'TEST','Camille',cid+'@example.test','MCO','presentiel',f'{year}BTS01120001',year,status,'2026-12-01T10:00:00',cid))


def payload():
    fields=dict(nom='TEST',prenom='Camille',sexe='Femme',date_naissance='2008-01-01',ville_naissance='Paris',cp_naissance='75001',pays_naissance='France',nationalite='Française',num_secu='208017512345646',email='camille@example.test',email_confirm='camille@example.test',tel='0600000000',adresse='1 rue Test',cp='75001',ville='Paris',bts='MCO',mode='presentiel',baccalaureat='Prévu en 2027',bac_type='Général',permis_b='Oui',projet_pourquoi='Apprendre le commerce',projet_objectif='Devenir responsable',projet_passions='La musique et les voyages',consent='1',entreprise_trouvee='non',recherches_commencees='oui',cherche_idf='non')
    for group in application.PROFILE_GROUPS:
        fields[group['name']+'[]'] = [group['options'][0][0]]
    for key in ['ci','ci2','carte_vitale','cv','lm']:
        fields[key]=(io.BytesIO(b'%PDF-1.4\n' + b' ' * 11000),key+'.pdf')
    fields['photo']=(io.BytesIO(b'\x89PNG\r\n\x1a\n'),'photo.png')
    return fields


def test_exact_paris_opening_boundary():
    assert not campaign_is_open(OPENS_AT-timedelta(microseconds=1))
    assert campaign_is_open(OPENS_AT)
    assert campaign_is_open(OPENS_AT+timedelta(seconds=1))
    assert OPENS_AT.isoformat()=='2026-12-01T09:00:00+00:00'


def test_closed_server_cannot_be_bypassed(client, monkeypatch):
    monkeypatch.setattr(application,'campaign_is_open',lambda:False)
    assert b'count-days' in client.get('/').data
    assert b'application-form' not in client.get('/?admin_bypass=1').data
    assert client.get('/candidature?preview=1').status_code==302
    assert client.post('/submit',data=payload()).status_code==403
    assert client.post('/save_draft',json={'full_form':{'email':'a@example.test'}}).status_code==403
    assert client.get('/reprendre/any-token').status_code==302
    auth(client)
    assert client.get('/candidature?preview=1').status_code==200
    assert client.post('/submit',data=payload()).status_code==403
    with application.db() as conn: assert conn.execute('SELECT COUNT(*) FROM candidats').fetchone()[0]==0


def test_opening_and_manual_pause(client,monkeypatch):
    monkeypatch.setattr(application,'campaign_is_open',lambda:True)
    assert client.get('/candidature').status_code==200
    assert client.post('/set_portal_status',json={'status':'closed'}).status_code==403
    auth(client)
    assert client.post('/set_portal_status',json={'status':'closed'}).status_code==200
    assert client.get('/candidature').status_code==302
    client.post('/set_portal_status',json={'status':'open'})
    assert client.get('/candidature').status_code==200


def test_ics_uses_absolute_timezone(client):
    response=client.get('/ouverture.ics')
    assert response.mimetype=='text/calendar'
    assert b'DTSTART:20261201T090000Z' in response.data


def test_trainer_pages_are_integrated_without_student_or_legacy_scripts(client, monkeypatch):
    monkeypatch.setattr(application.app, 'secret_key', 'integration-test-secret-for-trainers')
    login = client.get('/formateur/connexion')
    assert login.status_code == 200
    assert b'ESPACE FORMATEUR' in login.data
    assert b'trainer-space.css' in login.data
    for forbidden in (b'googletagmanager.com', b'cdn.brevo.com', b'js/main_front.js', b'js/main_admin.js'):
        assert forbidden not in login.data
    unsigned = client.get('/admin/formateurs?bts=MCO&rentree=2026')
    assert unsigned.status_code == 302
    assert unsigned.headers['Location'].startswith('/login?next=')
    assert 'formateurs' in unsigned.headers['Location']
    assert client.post('/admin/formateurs',json={}).status_code == 403
    auth(client)
    management = client.get('/admin/formateurs?bts=MCO&rentree=2026')
    assert management.status_code == 200
    assert b'id="workspace-sidebar"' in management.data
    assert b'/admin/formateurs' in management.data
    assert b'js/trainer-space.js' in management.data
    assert b'js/main_admin.js' not in management.data
    seed(2026, 'trainer-entry')
    fiche = client.get('/admin/candidat/trainer-entry?rentree=2026')
    assert fiche.status_code == 200
    assert b'/admin/formateurs?bts=MCO&amp;rentree=2026' in fiche.data


def test_explicit_promotion_on_student_sheet_and_course_change_guards(client):
    seed(2026, 'promotion-student')
    with application.db() as conn:
        conn.execute("UPDATE candidats SET tel='0600000000' WHERE id='promotion-student'")
    auth(client)
    assert client.get('/admin/promotions?bts=MCO&rentree=2026').status_code == 200
    with client.session_transaction() as state:
        token = state['_promotion_csrf']
    created = client.post('/admin/promotions', json=dict(csrf_token=token, nom='BTS MCO · 1re année', bts='MCO', niveau=1, rentree_annee=2026, date_debut='2026-09-07', date_fin='2027-06-30'))
    assert created.status_code == 201
    promotion_id = created.json['promotion_id']
    route = '/admin/candidat/promotion-student/promotion'
    assert client.post(route, json=dict(promotion_id=promotion_id)).status_code == 400
    assert client.post(route, json=dict(csrf_token=token, promotion_id=promotion_id)).json['ok']
    sheet = client.get('/admin/candidat/promotion-student')
    assert sheet.status_code == 200
    assert f'value="{promotion_id}" selected' in sheet.text
    assert f'/admin/formateurs?promotion={promotion_id}' in sheet.text
    assert '07/09/2026' in sheet.text and '30/06/2027' in sheet.text
    assert client.post('/admin/candidat/promotion-student/informations', json={'bts':'MOS'}).status_code == 409
    assert client.post('/admin/update-field', json={'id':'promotion-student','field':'bts','value':'MOS'}).status_code == 409
    with application.db() as conn:
        assert conn.execute("SELECT bts FROM candidats WHERE id='promotion-student'").fetchone()[0] == 'MCO'
    assert client.post(route, json=dict(csrf_token=token, promotion_id='')).json['ok']
    assert client.post('/admin/candidat/promotion-student/informations', json={'bts':'MOS'}).json['ok']
    assert client.post(route, json=dict(csrf_token=token, promotion_id=promotion_id)).status_code == 400


@pytest.mark.parametrize('clear_all', [False, True])
def test_candidate_deletion_removes_placement_but_keeps_the_promotion(client, clear_all):
    from promotions import assign_candidate_promotion
    seed(2026, 'removed-student')
    auth(client)
    client.get('/admin/promotions')
    with client.session_transaction() as state:
        token = state['_promotion_csrf']
    created = client.post('/admin/promotions', json=dict(csrf_token=token, nom='Promotion conservée', bts='MCO', niveau=1, rentree_annee=2026, date_debut='2026-09-07', date_fin='2027-06-30'))
    promotion_id = created.json['promotion_id']
    with application.db() as conn:
        assign_candidate_promotion(conn, 'removed-student', promotion_id)
    assert client.post('/admin/clear-db' if clear_all else '/admin/delete/removed-student').status_code in (200, 302)
    with application.db() as conn:
        assert conn.execute('SELECT COUNT(*) FROM student_promotions').fetchone()[0] == 0
        assert conn.execute('SELECT COUNT(*) FROM student_promotion_history').fetchone()[0] == 0
        assert conn.execute('SELECT COUNT(*) FROM promotions WHERE id=?', (promotion_id,)).fetchone()[0] == 1


def test_cohort_migration_idempotent_and_isolated_admin(client):
    seed(2026,'old');seed(2027,'new')
    application.ensure_campaign_schema();application.ensure_campaign_schema()
    auth(client)
    assert [r['id'] for r in client.get('/admin/json').json['rows']]==['new']
    assert [r['id'] for r in client.get('/admin/json?rentree=2026').json['rows']]==['old']
    assert len(client.get('/admin/json?rentree=all').json['rows'])==2
    assert client.get('/api/kpi?rentree=2027').json['confirmee']==1
    assert client.get('/admin/count_confirmed?rentree=2026').json['count']==1
    assert [r['id'] for r in client.get('/admin/export.json?rentree=2027').json]==['new']
    assert 'old' not in client.get('/admin/export.csv?rentree=2027').text


def test_complete_submission_saves_cohort_permit_and_both_identity_sides(client,monkeypatch):
    monkeypatch.setattr(application,'campaign_is_open',lambda:True)
    response=client.post('/submit',data=payload(),headers={'Accept':'application/json'})
    assert response.status_code==200, response.text
    assert response.json['ok']
    with application.db() as conn:
        row=dict(conn.execute('SELECT * FROM candidats').fetchone())
    assert row['rentree_annee']==2027
    assert row['numero_dossier'].startswith('2027BTS')
    assert row['permis_b']==1
    assert len(json.loads(row['fichiers_ci']))==2
    for f in json.loads(row['fichiers_ci']): assert os.path.isfile(f)
    page=client.get(response.json['redirect'])
    assert 'septembre 2027' in page.text
    assert 'septembre 2026' not in page.text


def test_incomplete_submission_does_not_write_files_or_candidates(client,monkeypatch):
    monkeypatch.setattr(application,'campaign_is_open',lambda:True)
    p=payload();p['email_confirm']='different@example.test'
    assert client.post('/submit',data=p).status_code==400
    with application.db() as conn: assert conn.execute('SELECT COUNT(*) FROM candidats').fetchone()[0]==0


def test_draft_resumption_retains_campaign_and_tags(client,monkeypatch):
    monkeypatch.setattr(application,'campaign_is_open',lambda:True)
    response=client.post('/save_draft',json={'full_form':{'email':'test@example.test','qualites[]':['rigoureux','curieux'],'bts':'CG'}})
    assert response.json['success']
    resumed=client.get(response.json['link'])
    assert resumed.status_code==200
    assert 'rigoureux' in resumed.text and 'curieux' in resumed.text
    assert 'rentree_annee' in resumed.text


def test_old_drafts_cannot_be_silently_reclassified(client,monkeypatch,tmp_path):
    monkeypatch.setattr(application,'campaign_is_open',lambda:True)
    monkeypatch.setenv('DATA_DIR',str(tmp_path))
    (tmp_path/'drafts.json').write_text(json.dumps([{'id':'old-draft','full_form':{'email':'old@example.test'}}]))
    response=client.get('/reprendre/old-draft')
    assert response.status_code==302 and '/candidature' in response.location


def test_mail_and_numbers_follow_candidate_year(client):
    assert dossier_number().startswith('2027BTS')
    assert candidate_year({'numero_dossier':'2026BTS01120001'})==2026
    with application.app.test_request_context():
        for year in [2026,2027]:
            context=get_mail_context({'rentree_annee':year,'bts':'MCO','prenom':'Camille','slug_public':'example'})
            html=mail_html('candidature_validee',**context)
            assert f'septembre {year}' in html
            assert f'septembre {2027 if year==2026 else 2026}' not in html


def test_ypareo_mos_never_reuses_previous_campaign(monkeypatch):
    from services.ypareo_neo import construire_payload_participation_bts_mos, YpareoError
    for key in ['ACTION_FORMATION_ID','DATE_DEBUT','DATE_FIN']:
        monkeypatch.delenv('YPAREO_BTS_MOS_2027_'+key, raising=False)
    with pytest.raises(YpareoError): construire_payload_participation_bts_mos(2027)
    assert construire_payload_participation_bts_mos(2026)[0]['dateDebut'].startswith('2026')
    monkeypatch.setenv('YPAREO_BTS_MOS_2027_ACTION_FORMATION_ID','123')
    monkeypatch.setenv('YPAREO_BTS_MOS_2027_DATE_DEBUT','2027-09-01')
    monkeypatch.setenv('YPAREO_BTS_MOS_2027_DATE_FIN','2028-06-30')
    item=construire_payload_participation_bts_mos(2027)[0]
    assert item['idActionFormation']==123 and item['dateDebut']=='2027-09-01'


def test_certificates_follow_each_candidates_cohort(client):
    from docx import Document
    seed(2026,'previous');seed(2027,'current');auth(client)
    for cid,year in [('previous',2026),('current',2027)]:
        for path in ['generate_certificat','generate_certificat_presentiel']:
            response=client.get(f'/admin/{path}/{cid}')
            assert response.status_code==200
            doc=Document(io.BytesIO(response.data))
            text=' '.join(p.text for p in doc.paragraphs)+' '.join(p.text for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs)
            assert str(year) in text and str(year+2) in text, text
            assert f"Certificat de scolarité {year}-{year+2}" in text
            assert '{{' not in text
            if year == 2027:
                assert '6000€' not in text
                assert 'septembre 2026' not in text
                assert '2025, 100 %' in text  # A historical result must remain historical.
            response.close()


def test_aps_2027_does_not_promise_unconfirmed_dates_or_prices(client):
    with application.app.test_request_context():
        html=mail_html('demande_aps',rentree_annee=2027,prenom='Camille')
        assert 'rentrée 2027' in html and 'confirmer le calendrier' in html
        assert '950' not in html and 'QONTO' not in html


def test_application_retains_training_guidance_and_next_steps(client, monkeypatch):
    monkeypatch.setattr(application, 'campaign_is_open', lambda: True)
    response = client.get('/candidature')
    assert response.status_code == 200
    html = response.text
    expected = {
        'MOS': ('41000', 'responsables d’équipes de sécurité privée'),
        'MCO': ('38362', 'management d’équipe dans tous les secteurs'),
        'NDRC': ('38368', 'la relation client, la vente et le marketing digital'),
        'PI': ('38380', 'conseillers immobiliers pour agences et syndics'),
        'CI': ('41759', 'import-export, prospection, négociation'),
        'CG': ('39159', 'la gestion comptable, du contrôle et de la finance'),
    }
    for code, (rncp, description) in expected.items():
        panel = html.split(f'data-training-info="{code}"', 1)[1].split('</article>', 1)[0]
        assert description in panel
        assert f'/recherche/rncp/{rncp}/' in panel
        assert '2 ans' in panel and 'Bac +2 · Niveau 5' in panel
        assert 'Examen officiel' in panel
    for phrase in ('ZOOM', 'emploi du temps fixe', 'évaluations et devoirs',
                   'académie de Nice', 'académie de Paris', 'CNAPS',
                   'autorisation préalable', 'Pôle Alternance',
                   '0 € pour l’apprenti', 'Aucun frais d’inscription',
                   'rémunéré selon ton âge', 'accusé de réception',
                   'numéro de dossier', 'Je confirme mon inscription'):
        assert phrase in html
    assert '8 juillet' not in html and '12 août 2026' not in html
    assert '15 décembre 2026' not in html and '950' not in html
    assert 'Sessions et tarif apprenti 2027' in html
    auth(client)
    assert 'data-training-info="CG"' in client.get('/candidature?preview=1').text


def test_certificate_placeholders_split_across_runs():
    from docx import Document
    from certificates import personalize_certificate
    doc=Document();p=doc.add_paragraph()
    p.add_run('2026-2028 — {{NOM_').bold=True
    p.add_run('PRENOM}} : {{DATE_AUJOURDHUI}}')
    personalize_certificate(doc,{'{{NOM_PRENOM}}':'Camille TEST','{{DATE_AUJOURDHUI}}':'24/10/2026'},2027)
    assert p.text=='2027-2029 — Camille TEST : 24/10/2026'
    assert p.runs[0].bold


def nir_with_key(core):
    numeric = core.replace('2A', '19').replace('2B', '18')
    return core + str(97 - int(numeric) % 97).zfill(2)


@pytest.mark.parametrize('core,birth,sex', [
    ('2080175123456','2008-01-01','Femme'),
    ('108012A123456','2008-01-01','Homme'),
    ('208012B123456','2008-01-01','Femme'),
    ('1080199123456','2008-01-01','Homme'),
])
def test_nir_key_corsica_and_birth_coherence(core,birth,sex):
    nir=nir_with_key(core)
    assert application.validate_nir(nir,birth,sex)[0]
    assert application.validate_nir(nir[:5]+' '+nir[5:],birth,sex)[0]
    assert not application.validate_nir(nir[:-2]+'00',birth,sex)[0]
    assert not application.validate_nir(nir,'2007-01-01',sex)[0]
    assert not application.validate_nir(nir,'2008-02-01',sex)[0]
    assert not application.validate_nir(nir,birth,'Homme' if sex=='Femme' else 'Femme')[0]
    assert not application.validate_nir(nir+'1',birth,sex)[0]


@pytest.mark.parametrize('name', ['qualites[]','motivation[]','valeurs[]','travail[]','projet_passions'])
def test_each_original_project_section_is_mandatory(client,monkeypatch,name):
    monkeypatch.setattr(application,'campaign_is_open',lambda:True)
    data=payload();data.pop(name)
    response=client.post('/submit',data=data)
    assert response.status_code==400, response.text
    with application.db() as conn:
        assert conn.execute('SELECT COUNT(*) FROM candidats').fetchone()[0]==0


def test_all_original_profile_choices_and_form_fields_remain(client):
    auth(client)
    html=client.get('/candidature?preview=1').text
    from html.parser import HTMLParser
    class Elements(HTMLParser):
        def __init__(self): super().__init__(); self.elements=[]
        def handle_starttag(self, tag, attrs): self.elements.append((tag,dict(attrs)))
    parser=Elements();parser.feed(html)
    expected={'qualites':['rigoureux','organise','autonome','communication','esprit_equipe','curieux','leadership','creatif','empathique','serieux'], 'motivation':['reussite','stabilite','utilite','defis','contact','innovation','evolution','apprentissage'], 'valeurs':['salaire','sens','contact','mobilite','stabilite','responsabilite','autonomie','equipe'], 'travail':['methodique','pratique','collaboratif','independant','creatif','objectif']}
    for group,values in expected.items():
        assert [attrs['value'] for tag,attrs in parser.elements if tag=='input' and attrs.get('name')==group+'[]']==values
    for name in ['bts','mode','nom','prenom','date_naissance','sexe','pays_naissance','ville_naissance','cp_naissance','nationalite','adresse','cp','ville','num_secu','email','email_confirm','tel','resp_nom','resp_prenom','resp_email','resp_tel','baccalaureat','bac_type','bac_autre','permis_b','entreprise_trouvee','recherches_commencees','cherche_idf','souhaite_accompagnement','projet_pourquoi','projet_objectif','projet_passions','raison_aps','aps_souhaitee','ci','ci2','photo','carte_vitale','cv','lm']:
        assert any(attrs.get('name')==name for tag,attrs in parser.elements), name
    assert sum('data-step' in attrs for tag,attrs in parser.elements)==7
    assert sum('data-suggest' in attrs for tag,attrs in parser.elements)==6


@pytest.mark.parametrize('modifications', [
    {'num_secu':'208017512345600'},
    {'bac_type':'Autre','bac_autre':''},
    {'bts':'MOS','mos_parcours':'autre','aps_souhaitee':'non','raison_aps':''},
    {'cherche_idf':'oui','souhaite_accompagnement':''},
    {'entreprise_trouvee':''},
    {'qualites[]':['unknown_value']},
    {'sexe':'unknown_value'},
])
def test_conditional_and_nir_validation_cannot_be_bypassed(client,monkeypatch,modifications):
    monkeypatch.setattr(application,'campaign_is_open',lambda:True)
    data=payload();data.update(modifications)
    assert client.post('/submit',data=data).status_code==400


@pytest.mark.parametrize('name,content,filename', [
    ('cv',b'%PDF-1.4\n','tiny.pdf'),
    ('cv',b'not a pdf'+b' '*11000,'fake.pdf'),
    ('cv',b'%PDF-1.4\n'+b' '*(8*1024*1024),'large.pdf'),
    ('photo',b'%PDF-1.4\n'+b' '*11000,'wrong.png'),
    ('photo',b'photo','wrong.exe'),
])
def test_file_controls_preserved(client,monkeypatch,name,content,filename):
    monkeypatch.setattr(application,'campaign_is_open',lambda:True)
    data=payload();data[name]=(io.BytesIO(content),filename)
    assert client.post('/submit',data=data).status_code==400
    assert not os.path.exists(application.UPLOAD_DIR) or not os.listdir(application.UPLOAD_DIR)


def test_minor_requires_valid_guardian_contacts(client,monkeypatch):
    monkeypatch.setattr(application,'campaign_is_open',lambda:True)
    from datetime import date
    year=date.today().year-17
    for invalid in [{}, {'resp_nom':'Parent','resp_prenom':'Camille','resp_email':'invalid','resp_tel':'0600000000'}]:
        data=payload();data.update(date_naissance=f'{year}-01-01',num_secu=nir_with_key(f'2{year%100:02d}0175123456'));data.update(invalid)
        assert client.post('/submit',data=data).status_code==400


def test_preview_never_creates_a_dossier_or_draft_even_after_opening(client,monkeypatch):
    auth(client);monkeypatch.setattr(application,'campaign_is_open',lambda:True)
    data=payload();data['admin_preview']='1'
    assert client.post('/submit',data=data).status_code==403
    response=client.post('/save_draft',json={'full_form':{'email':'test@example.test','admin_preview':'1'}})
    assert response.status_code==403
    assert not os.path.exists(os.path.join(application.DATA_DIR,'drafts.json'))


def test_candidate_sheet_retains_legacy_profiles_and_atomic_edits(client):
    seed(2026,'sheet')
    assert client.get('/admin/candidat/sheet').status_code==302
    assert client.post('/admin/candidat/sheet/informations',json={'prenom':'Alex'}).status_code==403
    with application.db() as conn:
        conn.execute("UPDATE candidats SET tel='0600000000',projet_qualites='rigoureux, curieux',projet_motivation='reussite',projet_recherche='sens',projet_travail='methodique' WHERE id='sheet'")
    auth(client)
    page=client.get('/admin/candidat/sheet');assert page.status_code==200
    for label in ['Fiche PDF','Supprimer le dossier','Rigueur','Curiosité','La réussite personnelle','Un métier qui a du sens','Je suis méthodique','panel-historique','panel-documents']:
        assert label in page.text
    for term in ['YPAREO','NEO','Intégration']:assert term not in page.text
    assert client.post('/admin/candidat/sheet/informations',json={'prenom':'Alex','email':'bad'}).status_code==400
    with application.db() as conn:assert conn.execute("SELECT prenom FROM candidats WHERE id='sheet'").fetchone()[0]=='Camille'
    assert client.post('/admin/candidat/sheet/informations',json={'prenom':'Alex','id':'unauthorized'}).json['ok']
    with application.db() as conn:assert conn.execute("SELECT prenom FROM candidats WHERE id='sheet'").fetchone()[0]=='Alex'
    logs=client.get('/admin/logs/sheet').json
    assert any(log['type']=='FICHE_MODIFIEE' for log in logs)
    assert client.get('/admin/candidat/missing').status_code==404
    assert client.post('/admin/update-status',json={'id':'sheet','value':'bad'}).status_code==400
    assert client.post('/admin/update-status',json={'id':'missing','value':'validee'}).status_code==404
    assert client.post('/admin/ypareo-neo/sheet').status_code==404
    assert client.post('/admin/update-field',json={'id':'sheet','field':'label_ypareo','value':True}).status_code==400


def test_country_city_address_and_prefix_suggestions_with_safe_fallback(client,monkeypatch):
    import place_search
    place_search.search_places.cache_clear();place_search.department_communes.cache_clear()
    assert any(v['value']=='États-Unis' for v in client.get('/api/admission/suggestions?kind=country&q=etats').json['items'])
    calls=[]
    def fake_get(url,params):
        calls.append((url,params))
        if 'geo.api.gouv.fr' in url:return [{'nom':'Puget-sur-Argens','codesPostaux':['83480'],'code':'83099'}]
        if 'photon' in url:return {'features':[{'properties':{'name':'Londres','country':'Royaume-Uni'}}]}
        return {'features':[{'properties':{'name':'54 Chemin du Carreou','label':'54 Chemin du Carreou 83480 Puget-sur-Argens','city':'Puget-sur-Argens','postcode':'83480','type':'housenumber'}}]}
    monkeypatch.setattr(place_search,'_get',fake_get)
    for query in ['kind=city&q=Puget','kind=postcode&q=83480','kind=postcode&q=834','kind=address&q=54+chemin']:
        data=client.get('/api/admission/suggestions?'+query).json['items'][0]
        assert data['city']=='Puget-sur-Argens' and data['postcode']=='83480'
    data=client.get('/api/admission/suggestions?kind=city&q=Londres&country=Royaume-Uni').json['items'][0]
    assert data['city']=='Londres'
    assert all(url.startswith('https://') for url,params in calls)
    import requests
    monkeypatch.setattr(place_search,'_get',lambda *a: (_ for _ in ()).throw(requests.Timeout()))
    assert client.get('/api/admission/suggestions?kind=city&q=Uncached').json['unavailable']
    assert client.get('/api/admission/suggestions?kind=bad&q=ab').status_code==400
    place_search.search_places.cache_clear();place_search.department_communes.cache_clear()


def test_pdf_zip_and_single_candidate_actions_remain_available(client,tmp_path):
    import zipfile
    seed(2027,'document-test');auth(client)
    document=tmp_path/'piece.pdf';document.write_bytes(b'%PDF-1.4\n'+b' '*11000)
    with application.db() as conn:
        conn.execute("UPDATE candidats SET fichiers_cv=?,tel='0600000000' WHERE id='document-test'",(json.dumps([str(document)]),))
    files=client.get('/admin/files/document-test').json
    assert files[0]['filename']=='piece.pdf'
    result=client.get('/admin/files/download/document-test')
    assert result.status_code==200
    assert zipfile.ZipFile(io.BytesIO(result.data)).namelist()==['piece.pdf'];result.close()
    result=client.get('/admin/print/document-test')
    assert result.status_code==200 and result.data.startswith(b'%PDF-'),result.text[:100];result.close()
    assert client.post('/admin/update-field',json={'id':'document-test','field':'label_aps','value':True}).json['ok']
    assert client.post('/admin/update-field',json={'id':'document-test','field':'commentaires','value':'Note de test'}).json['ok']
    assert client.post('/admin/delete/document-test').json['ok']
    assert client.get('/admin/candidat/document-test').status_code==404


def test_status_notification_is_preserved_without_duplicate_send(client,monkeypatch):
    seed(2027,'notification-test','preinscription');auth(client)
    mail=[];sms=[]
    monkeypatch.setattr(application,'send_mail',lambda *a,**k:mail.append(a))
    monkeypatch.setattr(application,'send_sms_brevo',lambda *a,**k:sms.append(a))
    with application.db() as conn:conn.execute("UPDATE candidats SET tel='0600000000' WHERE id='notification-test'")
    for _ in range(2):
        result=client.post('/admin/update-status',json={'id':'notification-test','value':'validee'})
        assert result.status_code==200
    assert len(mail)==1 and len(sms)==1
    assert 'septembre 2027' in mail[0][2]
    assert client.get('/admin/logs/notification-test').json


def test_identity_names_normalized_before_storage_and_draft(client, monkeypatch):
    monkeypatch.setattr(application, 'campaign_is_open', lambda: True)
    p = payload(); p['nom'] = "  d’arc  "; p['prenom'] = "  élISE  marie-clAIRE "
    response = client.post('/submit', data=p, headers={'Accept': 'application/json'})
    assert response.status_code == 200
    with application.db() as conn:
        row = conn.execute('SELECT nom,prenom FROM candidats').fetchone()
    assert tuple(row) == ('D’ARC', 'Élise Marie-Claire')
    assert application.normalize_candidate_names({'resp_nom': ' martin ', 'resp_prenom': 'jean-luc'}) == {'resp_nom': 'MARTIN', 'resp_prenom': 'Jean-Luc'}


def test_nir_wrong_body_with_original_key_does_not_blame_only_key():
    valid, _ = application.validate_nir('208017512345646', '2008-01-01', 'Femme')
    assert valid
    ok, message = application.validate_nir('208017512345746', '2008-01-01', 'Femme')
    assert not ok
    assert '15 caractères' in message and 'n’importe où' in message


def test_candidate_photo_is_scoped_and_private(client, tmp_path):
    from PIL import Image
    seed(2027, 'photo-owner'); seed(2027, 'another-person')
    folder = tmp_path / 'uploads' / 'photo-owner'; folder.mkdir(parents=True)
    path = folder / 'identity.png'; Image.new('RGB', (800, 900), 'gold').save(path)
    with application.db() as conn:
        conn.execute('UPDATE candidats SET fichiers_photo=? WHERE id=?', (json.dumps([str(path)]), 'photo-owner'))
        conn.execute('UPDATE candidats SET fichiers_photo=? WHERE id=?', (json.dumps([str(path)]), 'another-person'))
    assert client.get('/admin/candidat/photo-owner/photo').status_code == 403
    assert client.get('/espace/unknown/photo').status_code == 404
    assert client.get('/espace/another-person/photo').status_code == 404
    response = client.get('/espace/photo-owner/photo')
    assert response.status_code == 200 and response.mimetype == 'image/jpeg'
    assert response.headers['Cache-Control'] == 'private, no-store'
    with Image.open(io.BytesIO(response.data)) as image:
        assert max(image.size) <= 600
    auth(client)
    assert client.get('/admin/candidat/photo-owner/photo').status_code == 200


def test_login_code_email_does_not_copy_administration(monkeypatch):
    import utils
    sent = []
    class Response:
        ok = True
        text = '{"messageId":"test-only"}'
        def json(self): return {'messageId': 'test-only'}
    monkeypatch.setattr(utils, 'BREVO_KEY', 'test-only')
    monkeypatch.setattr(utils.requests, 'post', lambda *a, **kw: (sent.append(kw['json']) or Response()))
    utils.send_mail('candidate@example.test', 'Code test', '<p>Test</p>', copy_admin=False)
    assert sent[0]['to'] == [{'email': 'candidate@example.test'}]
    utils.send_mail('candidate@example.test', 'Dossier test', '<p>Test</p>')
    assert len(sent[1]['to']) == 2


def test_existing_signed_links_survive_session_key_configuration(client, monkeypatch):
    import hashlib, hmac
    seed(2026, 'legacy-link')
    token = 'old-candidate-confirmation-token-long-enough'
    with application.db() as conn:
        conn.execute('UPDATE candidats SET token_confirm=?,token_confirm_exp=? WHERE id=?', (token, '2099-01-01T00:00:00', 'legacy-link'))
    application.snapshot_legacy_links(application.db)
    old_signature = hmac.new(b'dev-secret', token.encode(), hashlib.sha256).hexdigest()
    monkeypatch.setenv('SECRET_KEY', 'new-unique-test-session-key')
    assert application.verify_token(token, old_signature)
    assert not application.verify_token('new-token-not-in-snapshot-long-enough', old_signature)
    assert not application.verify_token(token, 'é')


def test_old_photo_prefix_is_resolved_inside_own_dossier(client, tmp_path):
    from PIL import Image
    seed(2026, 'old-photo')
    folder = tmp_path / 'uploads' / 'old-photo'; folder.mkdir(parents=True)
    Image.new('RGB', (30, 40), 'gold').save(folder / 'PhotoIdentite_old-photo.jpg')
    with application.db() as conn:
        conn.execute('UPDATE candidats SET fichiers_photo=? WHERE id=?', (json.dumps(['/former/disk/uploads/old-photo/PhotoIdentite_old-photo.jpg']), 'old-photo'))
    assert client.get('/espace/old-photo/photo').status_code == 200


def test_legacy_flat_photo_belongs_to_exact_dossier(client, tmp_path):
    from PIL import Image
    seed(2026, 'flat-owner'); seed(2026, 'flat-other')
    folder = tmp_path / 'uploads'; folder.mkdir()
    photo = folder / 'Photo_flat-owner.jpeg'
    Image.new('RGB', (30, 40), 'gold').save(photo)
    with application.db() as conn:
        for cid in ('flat-owner', 'flat-other'):
            conn.execute('UPDATE candidats SET fichiers_photo=? WHERE id=?', (json.dumps([str(photo)]), cid))
    assert client.get('/espace/flat-owner/photo').status_code == 200
    assert client.get('/espace/flat-other/photo').status_code == 404


def test_jpeg_filename_containing_mpo_photo_is_supported(client, tmp_path):
    from PIL import Image
    seed(2026, 'mpo-owner')
    folder = tmp_path / 'uploads' / 'mpo-owner'; folder.mkdir(parents=True)
    photo = folder / 'Photo_mpo-owner.jpeg'
    Image.new('RGB', (30, 40), 'gold').save(photo, format='MPO', save_all=True, append_images=[Image.new('RGB', (30, 40), 'white')])
    with Image.open(photo) as image: assert image.format == 'MPO'
    with application.db() as conn:
        conn.execute('UPDATE candidats SET fichiers_photo=? WHERE id=?', (json.dumps([str(photo)]), 'mpo-owner'))
    response = client.get('/espace/mpo-owner/photo')
    assert response.status_code == 200 and response.mimetype == 'image/jpeg'


def test_student_portal_labels_and_password_links(client):
    seed(2027, 'student-preview')
    page = client.get('/espace/student-preview')
    assert page.status_code == 200
    assert 'ESPACE ÉTUDIANT' in page.text
    assert '/espace/mot-de-passe-oublie' in page.text
    assert 'espace candidat' not in page.text.lower()
    assert 'googletagmanager.com' not in page.text
    assert 'cdn.brevo.com' not in page.text


def test_admin_print_opens_inline_and_missing_person_is_404(client):
    seed(2026, 'print-inline')
    assert client.get('/admin/print/print-inline').status_code == 403
    auth(client)
    response = client.get('/admin/print/print-inline')
    assert response.status_code == 200
    assert response.mimetype == 'application/pdf'
    assert response.headers['Content-Disposition'].startswith('inline;')
    assert response.headers['Cache-Control'] == 'private, no-store'
    assert client.get('/admin/print/no-such-person').status_code == 404


def test_certificate_readable_preview_keeps_original_available(client):
    seed(2027, 'certificate-preview')
    auth(client)
    for route in ['generate_certificat', 'generate_certificat_presentiel']:
        response = client.get(f'/admin/{route}/certificate-preview?preview=1')
        assert response.status_code == 200
        assert response.mimetype == 'text/html'
        assert '2027' in response.text and '2029' in response.text
        assert 'Télécharger l’original' in response.text
        assert 'attachment' not in response.headers.get('Content-Disposition', '')


def test_old_upload_links_require_admin_and_forward_private_viewer(client, tmp_path):
    seed(2026, 'legacy-link')
    folder = tmp_path / 'uploads' / 'legacy-link'
    folder.mkdir(parents=True)
    (folder / 'proof.pdf').write_bytes(b'%PDF-1.4\n%%EOF')
    assert client.get('/uploads/proof.pdf').status_code == 403
    auth(client)
    response = client.get('/uploads/proof.pdf')
    assert response.status_code == 302
    assert response.location == '/admin/candidat/legacy-link/document/proof.pdf'
    response = client.get(response.location)
    assert response.headers['Content-Disposition'].startswith('inline;')
